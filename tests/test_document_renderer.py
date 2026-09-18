from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.schemas.generation import DifficultyCount, GenerationMode, QuestionType, QuestionTypeCount, VariationStrength
from app.schemas.papers import AddManualQuestionRequest, ExportFormat, ExportVariant, PaperCreateRequest, PaperUpdateRequest
from app.services.document_renderer import PaperDocumentRenderer, latex_to_readable
from app.services.papers import PaperService


class DocumentRendererTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        os.environ["DATABASE_URL"] = f"sqlite:///{Path(self.tmp.name) / 'test.db'}"
        request = PaperCreateRequest(
            title="Definite Integrals Practice Paper",
            exam="JEE Main",
            subject="Mathematics",
            chapters=["Calculus"],
            topics=["Definite Integrals"],
            question_types=[QuestionTypeCount(type=QuestionType.SINGLE_CORRECT, count=1)],
            difficulty_distribution=[DifficultyCount(difficulty=3, count=1)],
            generation_mode=GenerationMode.STRUCTURAL,
            variation_strength=VariationStrength.BALANCED,
        )
        self.paper = PaperService().create(request)
        self.paper = PaperService().update(
            self.paper["id"],
            PaperUpdateRequest(branding_config={
                "institution_name": "Apex Academy", "address": "21 Scholar Road, Pune", "contact": "hello@apex.example",
                "header_text": "Apex Academy · Mathematics", "footer_text": "Practice paper", "watermark_text": "APEX CONFIDENTIAL",
                "logo_data_url": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL2dAAAAABJRU5ErkJggg==",
                "duration_minutes": 60, "total_marks": 4,
            }),
        )
        PaperService().add_manual_question(
            self.paper["id"],
            AddManualQuestionRequest(
                question_type=QuestionType.SINGLE_CORRECT,
                stem="Evaluate $\\displaystyle \\int_0^1 x^2\\,dx$.",
                options=["$\\frac13$", "$\\frac12$", "$1$", "$0$"],
                correct_answer="A",
                solution="$\\int_0^1x^2\\,dx=\\frac13$",
                difficulty=3,
                marks=4,
                primary_concept="definite integral evaluation",
            ),
        )

    def test_latex_is_printable_unicode(self) -> None:
        rendered = latex_to_readable("$\\int_0^1 \\frac{(2r)^k}{n^{k+1}}\\,dx \\leq \\pi$")
        self.assertIn("∫₀¹", rendered)
        self.assertIn("≤ π", rendered)
        self.assertNotIn("frac", rendered)
        self.assertNotIn("frac", latex_to_readable("$\\frac1{k+1}$"))
        self.assertNotIn("frac", latex_to_readable("$\\frac1c\\int_a^b f\\left(\\frac{x}{c}\\right)dx$"))
        self.assertNotIn("\\", rendered)
        self.assertNotIn("$", rendered)

    def test_question_paper_and_answer_key_export(self) -> None:
        paper = PaperService().get(self.paper["id"])
        renderer = PaperDocumentRenderer(Path(self.tmp.name) / "exports")
        question_path, _ = renderer.export(paper, output_format=ExportFormat.DOCX, variant=ExportVariant.QUESTION_PAPER)
        answers_path, _ = renderer.export(paper, output_format=ExportFormat.DOCX, variant=ExportVariant.ANSWER_KEY)
        self.assertTrue(question_path.exists())
        self.assertTrue(answers_path.exists())
        question_document = Document(question_path)
        question_xml = question_document._element.xml
        question_table_text = "\n".join(cell.text for table in question_document.tables for row in table.rows for cell in row.cells)
        answers_text = "\n".join(paragraph.text for paragraph in Document(answers_path).paragraphs)
        self.assertNotIn("Candidate Name", question_table_text)
        self.assertNotIn("Roll Number", question_table_text)
        self.assertNotIn("JEE Main  |  Mathematics", "\n".join(paragraph.text for paragraph in question_document.paragraphs))
        self.assertIn("<m:oMath>", question_xml)
        self.assertIn("<m:sSubSup>", question_xml)
        self.assertIn("<m:t>∫</m:t>", question_xml)
        self.assertIn("<m:f>", question_xml)
        self.assertNotIn("\\frac", question_xml)
        self.assertNotIn("displaystyle", question_xml)
        self.assertIn("Answer Key", answers_text)
        self.assertEqual(question_document.styles["Normal"].font.name, "Times New Roman")
        self.assertIn("Apex Academy · Mathematics", "\n".join(item.text for item in question_document.sections[0].header.paragraphs))
        self.assertIn("Practice paper", "\n".join(item.text for item in question_document.sections[0].footer.paragraphs))
        self.assertAlmostEqual(question_document.sections[0].page_width, 7560000, delta=500)  # A4 width in EMU
        self.assertAlmostEqual(question_document.sections[0].page_height, 10692000, delta=500)  # A4 height in EMU

    def test_export_removes_xml_invalid_characters_from_generated_content(self) -> None:
        PaperService().add_manual_question(
            self.paper["id"],
            AddManualQuestionRequest(
                question_type=QuestionType.SINGLE_CORRECT,
                stem="A malformed model escape\x00 must not break $v=\x08$. ",
                options=["A", "B", "C", "D"],
                correct_answer="A",
                solution="Control\x0b character is removed.",
                difficulty=3,
                marks=1,
                primary_concept="safe rendering",
            ),
        )
        paper = PaperService().get(self.paper["id"])
        path, _ = PaperDocumentRenderer(Path(self.tmp.name) / "exports").export(
            paper, output_format=ExportFormat.DOCX, variant=ExportVariant.QUESTION_PAPER,
        )
        self.assertTrue(path.exists())
        rendered = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        self.assertNotIn("\x00", rendered)
        self.assertNotIn("\x08", rendered)
        self.assertNotIn("\x0b", rendered)

    def test_custom_layout_is_emitted_for_docx_and_pdf_source(self) -> None:
        paper = PaperService().get(self.paper["id"])
        paper["branding_config"]["layout"] = {
            "preset": "watermarked", "header_left": "Marks : 4", "header_center": "INTEGRALS TEST",
            "header_right": "Time : 60 min", "footer_center": "Apex practice", "watermark_enabled": True,
            "watermark_opacity": 16, "watermark_rotation": 315, "divider_enabled": True,
        }
        renderer = PaperDocumentRenderer(Path(self.tmp.name) / "exports")
        path, _ = renderer.export(paper, output_format=ExportFormat.DOCX, variant=ExportVariant.QUESTION_PAPER)
        document = Document(path)
        xml = document.sections[0].header._element.xml
        source = renderer._latex_source(paper, ExportVariant.QUESTION_PAPER)
        self.assertIn("INTEGRALS TEST", xml)
        self.assertIn("APEX CONFIDENTIAL", xml)
        self.assertIn(r"\fancyhead[L]", source)
        self.assertIn(r"\SetWatermarkText{APEX CONFIDENTIAL}", source)


if __name__ == "__main__":
    unittest.main()
